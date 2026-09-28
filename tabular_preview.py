"""Índice limitado de linhas normalizadas; leitura original sequencial uma única vez."""
import codecs
import csv
from datetime import date, datetime, time
from io import BytesIO, TextIOWrapper
from itertools import islice
from pathlib import Path
from zipfile import ZipFile

import xlrd
from fastapi import HTTPException
from openpyxl import load_workbook

from preview_cache import memory_size

PAGE_SIZE = 100
MAX_COLUMNS = 100
MAX_CELL_CHARS = 4000
MAX_INPUT_BYTES = 32 * 1024 * 1024
MAX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_INDEX_BYTES = 8 * 1024 * 1024
MAX_ROWS = 10_000
MAX_SHEETS = 64


def display_value(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, bool):
        return "VERDADEIRO" if value else "FALSO"
    return str(value)


def validate_table_input(content, filename):
    if len(content) > MAX_INPUT_BYTES:
        raise HTTPException(413, "Prévia limitada a arquivos de até 32 MiB. Baixe o original.")
    if Path(filename).suffix.lower() in (".xlsx", ".xlsm"):
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 2048 or sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise HTTPException(413, "Planilha muito grande após descompactação (limite 64 MiB). Baixe o original.")


def csv_encoding(content):
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    decoder = codecs.getincrementaldecoder("utf-8-sig")()
    try:
        for offset in range(0, len(content), 65536):
            decoder.decode(content[offset:offset + 65536], final=False)
        decoder.decode(b"", final=True)
        return "utf-8-sig"
    except UnicodeDecodeError:
        return "cp1252"


def build_table_index(content, filename, delimiter="auto"):
    workbook = stream = None
    extension = Path(filename).suffix.lower()
    try:
        validate_table_input(content, filename)
        notes = []
        if extension in (".csv", ".tsv"):
            encoding = csv_encoding(content)
            stream = TextIOWrapper(BytesIO(content), encoding=encoding, newline="")
            if delimiter == "auto":
                if extension == ".tsv":
                    delimiter = "\t"
                else:
                    try:
                        delimiter = csv.Sniffer().sniff(stream.read(65536), delimiters=",;\t|").delimiter
                    except csv.Error:
                        delimiter = ","
                    stream.seek(0)
            # Limite do parser por campo antes da normalização para 4.000 caracteres.
            csv.field_size_limit(1_000_000)
            reader = csv.reader(stream, delimiter=delimiter)
            sheets = [filename]
            def rows_for_sheet(index):
                return reader
            notes.append(f"Codificação: {encoding}. Separador: {delimiter!r}.")
        elif extension in (".xlsx", ".xlsm"):
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
            sheets = workbook.sheetnames[:MAX_SHEETS]
            def rows_for_sheet(index):
                worksheet = workbook[sheets[index]]
                return worksheet.iter_rows(values_only=True, max_col=min(worksheet.max_column or 1, MAX_COLUMNS + 1))
            notes.append("Fórmulas aparecem como texto, sem recálculo. Gráficos, imagens e estilos do Excel não são reproduzidos.")
            if len(workbook.sheetnames) > MAX_SHEETS:
                notes.append("A prévia mostra apenas as primeiras 64 abas.")
        elif extension == ".xls":
            workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
            sheets = workbook.sheet_names()[:MAX_SHEETS]
            def rows_for_sheet(index):
                worksheet = workbook.sheet_by_index(index)
                try:
                    for row in range(worksheet.nrows):
                        values = []
                        for column in range(min(worksheet.ncols, MAX_COLUMNS + 1)):
                            cell = worksheet.cell(row, column)
                            value = cell.value
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                value = xlrd.xldate_as_datetime(value, workbook.datemode)
                            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                                value = bool(value)
                            elif cell.ctype == xlrd.XL_CELL_ERROR:
                                value = xlrd.error_text_from_code.get(value, "#ERRO")
                            values.append(value)
                        yield values
                finally:
                    workbook.unload_sheet(index)
            notes.append("Exibe os valores salvos no XLS, sem recálculo de fórmulas ou reprodução de gráficos e estilos.")
            if workbook.nsheets > MAX_SHEETS:
                notes.append("A prévia mostra apenas as primeiras 64 abas.")
        else:
            raise HTTPException(415, "Formato de planilha não suportado")

        tables = []
        used = memory_size(sheets) + memory_size(notes)
        total_rows = 0
        for sheet in range(len(sheets)):
            data = []
            sheet_notes = list(notes)
            truncated_columns = truncated_cells = limited = False
            rows = rows_for_sheet(sheet) if total_rows < MAX_ROWS and used < MAX_INDEX_BYTES else iter(())
            try:
                for row in rows:
                    if total_rows >= MAX_ROWS:
                        limited = True
                        break
                    truncated_columns |= len(row) > MAX_COLUMNS
                    normalized = []
                    for value in islice(row, MAX_COLUMNS):
                        text = display_value(value)
                        if len(text) > MAX_CELL_CHARS:
                            truncated_cells = True
                            text = text[:MAX_CELL_CHARS] + "…"
                        normalized.append(text)
                    size = memory_size(normalized) + 16
                    if used + size > MAX_INDEX_BYTES:
                        limited = True
                        break
                    data.append(normalized)
                    used += size
                    total_rows += 1
            finally:
                if hasattr(rows, "close"):
                    rows.close()
            if truncated_columns:
                sheet_notes.append("A prévia mostra apenas as primeiras 100 colunas.")
            if truncated_cells:
                sheet_notes.append("Células longas foram limitadas a 4.000 caracteres na prévia.")
            if limited or total_rows >= MAX_ROWS or used >= MAX_INDEX_BYTES:
                sheet_notes.append("Prévia parcial: limite de 10.000 linhas ou 8 MiB de índice por arquivo. Baixe o original para ver tudo.")
            tables.append({"rows": data, "notes": sheet_notes})
        return {"sheets": sheets, "tables": tables}
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(422, "Não foi possível ler a planilha. Verifique o formato, a codificação ou a proteção por senha.") from error
    finally:
        if stream is not None:
            stream.close()
        if workbook is not None:
            if extension == ".xls":
                workbook.release_resources()
            else:
                workbook.close()


def table_page(index, sheet=0, offset=0, header=False):
    if not 0 <= sheet < len(index["sheets"]):
        raise HTTPException(404, "Aba não encontrada")
    table = index["tables"][sheet]
    rows = table["rows"]
    start = offset + int(header)
    return {"sheets": index["sheets"], "sheet": sheet, "offset": offset,
            "rows": rows[start:start + PAGE_SIZE],
            "headers": (rows[0] if rows else []) if header else None,
            "has_more": len(rows) > start + PAGE_SIZE, "page_size": PAGE_SIZE,
            "first_row": start + 1, "notes": table["notes"]}


def read_table(content, filename, sheet=0, offset=0, delimiter="auto", header=False):
    return table_page(build_table_index(content, filename, delimiter), sheet, offset, header)
